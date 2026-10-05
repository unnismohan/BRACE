"""BRACE routes.files — extracted application responsibility."""

import asyncio
import io
import os
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Optional
from fastapi import Depends, File, HTTPException, UploadFile
from fastapi.responses import PlainTextResponse, StreamingResponse
from pydantic import BaseModel
from db import get_db
import authentication as mod_authentication
from fastapi import APIRouter

router = APIRouter()


@router.get("/api/projects/{project_id}/base-path")
def get_base_path(project_id: int, user=Depends(mod_authentication._proj_viewer)):
    """Container-absolute path to this project's suites dir (for locating a file on disk/PVC)."""
    import runtime as mod_runtime

    return {"base_path": str(mod_runtime._project_suites(project_id).resolve())}


@router.get("/api/projects/{project_id}/files")
def list_files(project_id: int, user=Depends(mod_authentication._proj_viewer)):
    import runtime as mod_runtime

    suites_dir = mod_runtime._project_suites(project_id)
    if not suites_dir.exists():
        return {}
    tree = {}
    base = str(suites_dir)
    for dirpath, dirnames, filenames in os.walk(base):
        dirnames.sort()
        rel_dir = os.path.relpath(dirpath, base).replace("\\", "/")
        folder = "" if rel_dir == "." else rel_dir
        for name in sorted(filenames):
            path = folder + "/" + name if folder else name
            tree.setdefault(folder, []).append({"name": name, "path": path})
    return tree


@router.get("/api/projects/{project_id}/files/download-all")
async def download_all_files(
    project_id: int, user=Depends(mod_authentication._proj_viewer)
):
    """Zip the whole project scripts tree and stream it to the browser."""
    import runtime as mod_runtime

    suites_dir = mod_runtime._project_suites(project_id)
    if not suites_dir.exists():
        raise HTTPException(404, "No scripts to export")
    project = (
        get_db()
        .execute("SELECT name FROM projects WHERE id=?", (project_id,))
        .fetchone()
    )
    proj_name = re.sub(
        "[^A-Za-z0-9._-]+", "_", project["name"] if project else str(project_id)
    )
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"brace-{proj_name}-scripts-{stamp}.zip"

    def build_zip() -> io.BytesIO:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for path in suites_dir.rglob("*"):
                if path.is_file():
                    zf.write(path, path.relative_to(suites_dir))
        buf.seek(0)
        return buf

    buf = await asyncio.to_thread(build_zip)

    def stream():
        while chunk := buf.read(65536):
            yield chunk

    return StreamingResponse(
        stream(),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/api/projects/{project_id}/excel/{filepath:path}")
def read_excel(
    project_id: int, filepath: str, user=Depends(mod_authentication._proj_viewer)
):
    import runtime as mod_runtime
    import csv, io
    import openpyxl

    path = mod_runtime._safe_path(project_id, filepath)
    if not path.exists():
        raise HTTPException(404, "File not found")
    ext = path.suffix.lower()
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = {}
        ROW_LIMIT = 2000
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if c is None else str(c) for c in row])
                if len(rows) >= ROW_LIMIT:
                    break
            sheets[name] = rows
        wb.close()
        return {"type": "xlsx", "sheets": sheets, "sheet_names": list(wb.sheetnames)}
    elif ext == ".csv":
        text = path.read_text(errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows = [list(r) for r in reader]
        return {"type": "csv", "sheets": {"Sheet1": rows}, "sheet_names": ["Sheet1"]}
    raise HTTPException(400, "Not an Excel/CSV file")


@router.put("/api/projects/{project_id}/excel/{filepath:path}")
def write_excel(
    project_id: int,
    filepath: str,
    body: dict,
    user=Depends(mod_authentication._proj_tester),
):
    import runtime as mod_runtime
    import csv, io
    import openpyxl

    path = mod_runtime._safe_path(project_id, filepath)
    ext = path.suffix.lower()
    sheets: dict = body.get("sheets", {})
    path.parent.mkdir(parents=True, exist_ok=True)
    if ext == ".xlsx":
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet_name, rows in sheets.items():
            ws = wb.create_sheet(sheet_name)
            for row in rows:
                ws.append(row if row else [""])
        wb.save(path)
    elif ext == ".csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        rows = next(iter(sheets.values()), [])
        for row in rows:
            writer.writerow(row)
        path.write_text(buf.getvalue())
    return {"saved": filepath}


@router.get("/api/projects/{project_id}/files/{filepath:path}")
def read_file(
    project_id: int, filepath: str, user=Depends(mod_authentication._proj_viewer)
):
    import runtime as mod_runtime

    path = mod_runtime._safe_path(project_id, filepath)
    if not path.exists() or not path.is_file():
        raise HTTPException(404, "File not found")
    return PlainTextResponse(path.read_text(errors="replace"))


@router.put("/api/projects/{project_id}/files/{filepath:path}")
def save_file(
    project_id: int,
    filepath: str,
    body: dict,
    user=Depends(mod_authentication._proj_tester),
):
    import runtime as mod_runtime

    path = mod_runtime._safe_path(project_id, filepath)
    if path.suffix not in mod_runtime.ALLOWED_EXTS:
        raise HTTPException(400, f"Extension {path.suffix} not allowed")
    path.parent.mkdir(parents=True, exist_ok=True)
    existed = path.exists()
    path.write_text(body.get("content", ""))
    mod_runtime.audit(
        user,
        "script.save",
        project_id=project_id,
        target=filepath,
        created=not existed,
        bytes=len(body.get("content", "")),
    )
    return {"saved": filepath}


@router.post("/api/projects/{project_id}/files/upload")
async def upload_files(
    project_id: int,
    files: list[UploadFile] = File(...),
    user=Depends(mod_authentication._proj_tester),
):
    import runtime as mod_runtime

    suites_dir = mod_runtime._project_suites(project_id)
    suites_dir.mkdir(parents=True, exist_ok=True)
    saved = []
    for uf in files:
        filename = Path(uf.filename).name
        if not filename:
            continue
        data = await uf.read()
        if filename.endswith(".zip"):
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                for member in zf.infolist():
                    if member.is_dir():
                        continue
                    mp = Path(member.filename)
                    if mp.suffix not in mod_runtime.ALLOWED_EXTS or ".git" in str(mp):
                        continue
                    dest = (suites_dir / mp).resolve()
                    if not mod_runtime._contained(suites_dir.resolve(), dest):
                        continue
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    dest.write_bytes(zf.read(member))
                    saved.append(str(mp))
        elif Path(filename).suffix in mod_runtime.ALLOWED_EXTS:
            dest = suites_dir / filename
            dest.write_bytes(data)
            saved.append(filename)
    mod_runtime.audit(user, "script.upload", project_id=project_id, files=len(saved))
    return {"uploaded": saved}


@router.delete("/api/projects/{project_id}/files/{filepath:path}")
def delete_file(
    project_id: int, filepath: str, user=Depends(mod_authentication._proj_tester)
):
    import runtime as mod_runtime

    path = mod_runtime._safe_path(project_id, filepath)
    if not path.exists():
        raise HTTPException(404, "File not found")
    path.unlink()
    mod_runtime.audit(user, "script.delete", project_id=project_id, target=filepath)
    return {"deleted": filepath}


class RenameReq(BaseModel):
    old_path: str
    new_path: str


@router.post("/api/projects/{project_id}/fs/rename")
def fs_rename(
    project_id: int, req: RenameReq, user=Depends(mod_authentication._proj_tester)
):
    """Rename or move a file/folder within the project suites dir."""
    import runtime as mod_runtime

    src = mod_runtime._safe_path(project_id, req.old_path)
    dst = mod_runtime._safe_path(project_id, req.new_path)
    if not src.exists():
        raise HTTPException(404, "Source not found")
    if dst.exists():
        raise HTTPException(409, f"'{req.new_path}' already exists")
    if src.is_file() and dst.suffix not in mod_runtime.ALLOWED_EXTS:
        raise HTTPException(400, f"Extension {dst.suffix} not allowed")
    dst.parent.mkdir(parents=True, exist_ok=True)
    src.rename(dst)
    mod_runtime.audit(
        user,
        "script.rename",
        project_id=project_id,
        target=req.new_path,
        **{"from": req.old_path},
    )
    return {"renamed": req.new_path}


class MkdirReq(BaseModel):
    path: str


@router.post("/api/projects/{project_id}/fs/mkdir")
def fs_mkdir(
    project_id: int, req: MkdirReq, user=Depends(mod_authentication._proj_tester)
):
    import runtime as mod_runtime

    path = mod_runtime._safe_path(project_id, req.path)
    if path.exists():
        raise HTTPException(409, "Already exists")
    path.mkdir(parents=True)
    (path / ".gitkeep").write_text("")
    return {"created": req.path}


@router.delete("/api/projects/{project_id}/fs/rmdir/{dirpath:path}")
def fs_rmdir(
    project_id: int, dirpath: str, user=Depends(mod_authentication._proj_tester)
):
    """Recursively delete a folder. Refuses the project root."""
    import runtime as mod_runtime

    path = mod_runtime._safe_path(project_id, dirpath)
    base = mod_runtime._project_suites(project_id).resolve()
    if path == base:
        raise HTTPException(400, "Cannot delete the project root")
    if not path.exists() or not path.is_dir():
        raise HTTPException(404, "Folder not found")
    n = sum((1 for _ in path.rglob("*") if _.is_file()))
    shutil.rmtree(path)
    mod_runtime.audit(
        user, "script.rmdir", project_id=project_id, target=dirpath, files_removed=n
    )
    return {"deleted": dirpath, "files_removed": n}


@router.get("/api/projects/{project_id}/suites")
def list_suites(project_id: int, user=Depends(mod_authentication._proj_viewer)):
    """Return list of .robot file paths relative to project suites dir."""
    import runtime as mod_runtime

    suites_dir = mod_runtime._project_suites(project_id)
    if not suites_dir.exists():
        return []
    results = []
    for p in sorted(suites_dir.rglob("*.robot")):
        rel = str(p.relative_to(suites_dir)).replace("\\", "/")
        parts = rel.lower().split("/")
        if any((part == "testcases" for part in parts[:-1])):
            results.append(rel)
    return results


class QuickRunReq(BaseModel):
    suite_path: str
    extra_args: Optional[str] = None


@router.post("/api/projects/{project_id}/quick-run")
async def quick_run(
    project_id: int, req: QuickRunReq, user=Depends(mod_authentication._proj_tester)
):
    import runtime as mod_runtime

    suites_dir = mod_runtime._project_suites(project_id)
    target = mod_runtime._safe_path(project_id, req.suite_path)
    if not target.exists():
        raise HTTPException(404, "File not found")
    run_id = f"qr-{project_id}-{int(datetime.now().timestamp() * 1000)}"
    run_dir = mod_runtime._project_results(project_id) / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    mod_runtime.audit(user, "run.quick", project_id=project_id, target=req.suite_path)
    cmd = [
        "python",
        "-m",
        "robot",
        "--outputdir",
        str(run_dir),
        "--output",
        "output.xml",
        "--log",
        "log.html",
        "--report",
        "report.html",
        "--pythonpath",
        str(suites_dir),
        "--variable",
        f"BSS_ENV:{mod_runtime.BSS_ENV}",
    ]
    if req.extra_args:
        cmd.extend(mod_runtime._split_args(req.extra_args))
    cmd.append(str(target))

    async def stream():
        import runtime as mod_runtime

        yield f"[BRACE] Quick run: {req.suite_path}\n[BRACE] CMD: {' '.join(cmd)}\n\n"
        if mod_runtime._tslots().locked():
            yield "[BRACE] All execution slots busy — waiting for one to free up…\n"
        async with mod_runtime._tslots():
            env = {**os.environ, "DISPLAY": ":99"}
            proc = await asyncio.create_subprocess_exec(
                *cmd,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.STDOUT,
                env=env,
            )
            async for line in proc.stdout:
                yield line.decode(errors="replace")
            await proc.wait()
            rc = proc.returncode
        yield f"\n[BRACE] Exit code: {rc} — {('PASS' if rc == 0 else 'FAIL')}\n"
        yield f"[BRACE] Report: /results/{project_id}/{run_id}/report.html\n"

    return StreamingResponse(stream(), media_type="text/plain")
